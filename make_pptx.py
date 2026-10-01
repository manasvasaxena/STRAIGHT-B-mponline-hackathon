import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette (Dark Theme / High Tech Global Hackathon Vibe)
    BG_COLOR = RGBColor(11, 15, 25)       # Deep Obsidian
    CARD_BG = RGBColor(21, 29, 46)        # Dark Navy
    CARD_BORDER = RGBColor(45, 64, 97)    # Muted Blue Border
    ACCENT_CYAN = RGBColor(6, 182, 212)   # Bright Cyan
    ACCENT_BLUE = RGBColor(59, 130, 246)  # Electric Blue
    ACCENT_GREEN = RGBColor(16, 185, 129) # Emerald Green
    TEXT_WHITE = RGBColor(255, 255, 255)
    TEXT_MUTED = RGBColor(156, 163, 175)
    TEXT_SUBTITLE = RGBColor(203, 213, 225)

    def set_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title, category):
        # Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN
        p_cat.font.name = "Calibri"

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.733), Inches(0.6))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.font.name = "Calibri"

    def add_card(slide, left, top, width, height, title, accent_color, items):
        # Card Background Box
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        # Header Box inside Card
        header_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.5))
        tf_h = header_box.text_frame
        tf_h.word_wrap = True
        p_h = tf_h.paragraphs[0]
        p_h.text = title
        p_h.font.size = Pt(16)
        p_h.font.bold = True
        p_h.font.color.rgb = accent_color
        p_h.font.name = "Calibri"

        # Content Box
        content_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.65), width - Inches(0.4), height - Inches(0.8))
        tf_c = content_box.text_frame
        tf_c.word_wrap = True

        for idx, (head, desc) in enumerate(items):
            p = tf_c.add_paragraph() if idx > 0 else tf_c.paragraphs[0]
            p.space_after = Pt(8)
            
            run_head = p.add_run()
            run_head.text = head + "\n"
            run_head.font.bold = True
            run_head.font.size = Pt(12)
            run_head.font.color.rgb = TEXT_WHITE
            run_head.font.name = "Calibri"

            run_desc = p.add_run()
            run_desc.text = desc
            run_desc.font.size = Pt(10.5)
            run_desc.font.color.rgb = TEXT_MUTED
            run_desc.font.name = "Calibri"

    # ==================== SLIDE 1: TECH STACK & ARCHITECTURE ====================
    slide1 = prs.slides.add_slide(blank_layout)
    set_background(slide1)
    add_header(slide1, "SUTRA — Full-Stack Deep-Tech Architecture (Android & Web)", "Global Hackathon Deck | Slide 1 of 2")

    col_width = Inches(3.64)
    gap = Inches(0.4)
    top_pos = Inches(1.5)
    card_h = Inches(5.4)

    # Col 1: Native Android (APK)
    items_android = [
        ("UI & Modern Jetpack Compose", "Declarative UI rendered via Material 3, optimized for high-performance low-spec Android devices (Min SDK 24 / Target SDK 37)."),
        ("Zero-Internet P2P Battle Engine", "Native Wi-Fi Direct (P2P Manager) & TCP/UDP Sockets enable instant multi-device multiplayer quiz battles with zero internet."),
        ("Offline Markdown & Package Parser", "Custom CommonMark GFM parser renders local PKGS bundles, markdown notes, and tables with lightning-fast native speeds."),
        ("Resilient State Architecture", "Unidirectional Data Flow with Kotlin StateFlow & Coroutines, ensuring zero UI drops during heavy peer networking.")
    ]
    add_card(slide1, Inches(0.8), top_pos, col_width, card_h, "Native Android App (APK)", ACCENT_CYAN, items_android)

    # Col 2: Web Application Platform
    items_web = [
        ("Lightweight Vanilla Web Stack", "High-efficiency HTML5, CSS3, and ES6 JavaScript engine built for maximum speed and instant loading across any browser."),
        ("Offline LocalStorage Persistence", "Full client-side state machine caching quiz progress, bookmarks, and user session metrics directly on device."),
        ("Cross-Platform Web Engine", "Unified component structure guaranteeing 100% feature parity with Android app UI across tablets, PCs, and web kiosks."),
        ("Zero-Dependency Execution", "Pure native browser capabilities eliminate external bundle overhead, ensuring instant boot times even on legacy hardware.")
    ]
    add_card(slide1, Inches(0.8) + col_width + gap, top_pos, col_width, card_h, "Web Application Engine", ACCENT_BLUE, items_web)

    # Col 3: Local PKGS & Synchronisation Core
    items_core = [
        ("Local PKGS Content System", "Pre-packaged offline learning bundles (MD format) providing instant subject access without cloud round-trips."),
        ("Sync Queue & Offline Buffer", "Intelligent transaction logger queues user progress offline and seamlessly reconciles states upon network availability."),
        ("Cross-Layer Protocol Contract", "Shared JSON/Markdown schemas ensure frictionless data compatibility between Android native and Web platforms."),
        ("Production Build System", "Gradle build pipelines configured for dynamic artifact naming (`sutra.apk`) and automated optimized web bundling.")
    ]
    add_card(slide1, Inches(0.8) + (col_width + gap) * 2, top_pos, col_width, card_h, "Offline Core & Sync Engine", ACCENT_GREEN, items_core)

    # ==================== SLIDE 2: END-TO-END WORKFLOW & DATA LIFECYCLE ====================
    slide2 = prs.slides.add_slide(blank_layout)
    set_background(slide2)
    add_header(slide2, "End-to-End System Workflow & Peer-to-Peer Data Lifecycle", "Global Hackathon Deck | Slide 2 of 2")

    col2_w = Inches(5.66)
    gap2 = Inches(0.4)

    # Workflow 1: Content Ingestion to Render & Sync
    items_wf1 = [
        ("1. Offline PKGS Content Ingestion", "Local Markdown packages (.md) loaded instantly from local storage into memory; CommonMark parses GFM tables & formatting."),
        ("2. Interactive Student Execution", "Students engage in interactive lessons, offline doubt solving, and structured quizzes across Android or Web interface."),
        ("3. Atomic Local Storage & Buffering", "All scores, quiz completion metrics, and state changes are written instantly to local state repository and sync queue."),
        ("4. Seamless Background Synchronization", "When connectivity is detected, Sync Queue automatically flushes buffered transactions to server without user friction.")
    ]
    add_card(slide2, Inches(0.8), top_pos, col2_w, card_h, "A. Offline Core & Data Lifecycle Workflow", ACCENT_CYAN, items_wf1)

    # Workflow 2: Zero-Internet P2P Battle Workflow
    items_wf2 = [
        ("1. Wi-Fi Direct Peer Discovery", "Host initiates battle room; WifiDirectManager scans local frequency bands and establishes secure direct connection."),
        ("2. Socket Handshake & Room Setup", "TCP Server Socket opens on Host; Client connects via IP. Question packages and battle configuration synced in <50ms."),
        ("3. Real-Time Peer Battle Mechanics", "Players submit answers simultaneously; binary battle protocol broadcasts live score updates and timing over local sockets."),
        ("4. Winner Reconciliation & Broadcast", "Host calculates final leaderboards, generates BattleFinalSummary, and broadcasts final results to all connected peers.")
    ]
    add_card(slide2, Inches(0.8) + col2_w + gap2, top_pos, col2_w, card_h, "B. P2P Quiz Battle Real-Time Workflow", ACCENT_GREEN, items_wf2)

    # Save presentation
    output_path = "SUTRA_Hackathon_Master_Deck.pptx"
    prs.save(output_path)
    print(f"Successfully generated presentation at: {output_path}")

if __name__ == "__main__":
    create_deck()
